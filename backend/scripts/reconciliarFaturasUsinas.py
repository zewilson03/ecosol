"""Reconcilia UCs de faturas pendentes com planilhas, sem alterar as fontes.

Uso: python reconciliarFaturasUsinas.py PASTA_PLANILHAS BANCO [--apply HASH]
Sem --apply imprime somente a previa. A aplicacao exige o hash da previa,
faz backup consistente e aborta integralmente em caso de divergencia.
"""

import argparse
import hashlib
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from importarDescontosUcsPlanilhasUsinas import percentual
from vincularClientesPorPlanilhasUsinas import ARQUIVOS, nome_valido, normalizar, uc


APROVADA = "Aprovada — aguardando emissão"
MESES = {"2026-09": "set", "2026-10": "out"}


def hash_arquivo(caminho):
    h = hashlib.sha256()
    with caminho.open("rb") as fonte:
        for parte in iter(lambda: fonte.read(1024 * 1024), b""):
            h.update(parte)
    return h.hexdigest()


def mes_aba(titulo):
    texto = normalizar(titulo)
    if re.search(r"\bout(ubro)?\b", texto):
        return "2026-10"
    if re.search(r"\bset(embro)?\b", texto):
        return "2026-09"
    return None


def chave_usina(nome):
    return re.sub(r"[^a-z0-9]", "", normalizar(nome))


def fontes(pasta):
    """Consulta primeiro o mapa geral; depois usa abas recentes das usinas."""
    mapa_geral = defaultdict(set)
    geral = pasta / "uc por usina.xlsx"
    hashes = {geral.name: hash_arquivo(geral)}
    livro = load_workbook(geral, read_only=True, data_only=True)
    try:
        for numero, linha in enumerate(livro["Consulta"].iter_rows(min_row=2, values_only=True), 2):
            if len(linha) < 4:
                continue
            for valor in linha[:2]:
                if uc(valor):
                    mapa_geral[uc(valor)].add((str(linha[3] or "").strip(), numero))
    finally:
        livro.close()
    linhas = defaultdict(lambda: defaultdict(list))
    for usina, arquivo in ARQUIVOS.items():
        caminho = pasta / arquivo
        hashes[arquivo] = hash_arquivo(caminho)
        livro = load_workbook(caminho, read_only=True, data_only=True)
        try:
            for aba in livro.worksheets:
                mes = mes_aba(aba.title)
                if not mes:
                    continue
                cabecalho = [normalizar(x) for x in next(aba.iter_rows(max_row=1, values_only=True))]
                if cabecalho[1] not in {"cliente", "clientes"}:
                    raise ValueError(f"Cabeçalho de cliente desconhecido: {arquivo}/{aba.title}")
                indices = [i for i, x in enumerate(cabecalho) if "uc" in x and
                           ("uc2" in x or "nova" in x or "atual" in x)]
                if len(indices) != 1:
                    raise ValueError(f"Coluna de UC atual ambígua: {arquivo}/{aba.title}")
                for numero, linha in enumerate(aba.iter_rows(min_row=2, values_only=True), 2):
                    identificador = uc(linha[indices[0]] if len(linha) > indices[0] else None)
                    if not identificador:
                        continue
                    linhas[identificador][mes].append({
                        "usina": usina,
                        "nome": nome_valido(linha[1]),
                        "uc_antiga": uc(linha[2]) if len(linha) > 2 else "",
                        "desconto": percentual(linha[6] if len(linha) > 6 else None),
                        "arquivo": arquivo, "aba": aba.title, "linha": numero,
                    })
        finally:
            livro.close()
        if hash_arquivo(caminho) != hashes[arquivo]:
            raise ValueError(f"Planilha mudou durante a leitura: {arquivo}")
    if hash_arquivo(geral) != hashes[geral.name]:
        raise ValueError("Planilha UC por usina mudou durante a leitura")
    return mapa_geral, linhas, hashes


def planejar(banco, mapa, planilhas, hashes):
    usinas = {r["nome"]: r["id"] for r in banco.execute("SELECT id,nome FROM usinas")}
    modalidades = {r["usina_id"]: r["modalidade"] for r in banco.execute(
        "SELECT usina_id,modalidade FROM usinas_modalidades WHERE inicio='2026-10'")}
    unidades = {uc(r["uc"]): dict(r) for r in banco.execute(
        "SELECT id,uc,uc_antiga,nome,cliente_id,cliente_provisorio_id,usina_id,versao FROM faturamento_unidades")}
    documentos = list(banco.execute("""
        SELECT id,extracao,dados,status,unidade_id,competencia FROM faturamento_documentos
        WHERE classificacao='fatura' AND excluido_em IS NULL ORDER BY id
    """))
    acoes, pendencias = [], []
    for documento in documentos:
        numero = uc(json.loads(documento["extracao"]).get("dados", {}).get("uc"))
        mes_fatura = json.loads(documento["dados"]).get("competencia") or documento["competencia"]
        unidade = unidades.get(numero)
        if documento["status"] == APROVADA:
            continue
        por_mes = {}
        motivos = []
        for mes in MESES:
            candidatas = planilhas.get(numero, {}).get(mes, [])
            assinaturas = {(x["usina"], normalizar(x["nome"]), x["desconto"]) for x in candidatas}
            if len(assinaturas) > 1:
                motivos.append(f"divergência na planilha da usina em {mes}")
            elif candidatas:
                por_mes[mes] = candidatas[0]
        usinas_encontradas = {x["usina"] for x in por_mes.values()}
        if len(usinas_encontradas) > 1:
            motivos.append("mudança de usina entre setembro e outubro")
        principal = por_mes.get("2026-10") or por_mes.get("2026-09")
        if not principal:
            motivos.append("UC ausente nas planilhas das usinas ativas")
        elif not principal["nome"]:
            motivos.append("nome ausente na planilha")
        geral = sorted(mapa.get(numero, []))
        nomes_mapa = {chave_usina(x[0]) for x in geral}
        # O mapa geral pode estar atrasado ou indicar a fila provisória Ecosol.
        # Uma divergência com outra usina ativa exige conferência humana.
        confirmadas = {chave_usina(x) for x in usinas if chave_usina(x) != "ecosol"}
        if principal and any(x in confirmadas and x != chave_usina(principal["usina"])
                             for x in nomes_mapa):
            motivos.append("usina do mapa geral difere da planilha mensal")
        if unidade and unidade["usina_id"] not in (None, usinas.get("ecosol"),
                                                    usinas.get(principal["usina"]) if principal else None):
            motivos.append("UC já pertence a outra usina no Ecosol")
        if unidade and banco.execute("SELECT 1 FROM faturamento_vinculos_usina WHERE unidade_id=? LIMIT 1",
                                    (unidade["id"],)).fetchone():
            motivos.append("UC já possui histórico de usinas; conferir antes de modificar")
        if unidade and banco.execute("SELECT 1 FROM faturamento_documentos WHERE unidade_id=? AND status=? LIMIT 1",
                                    (unidade["id"], APROVADA)).fetchone():
            motivos.append("UC possui cobrança aprovada")
        if unidade:
            for mes, fonte in por_mes.items():
                desconto_atual = banco.execute(
                    "SELECT desconto FROM faturamento_descontos_uc WHERE unidade_id=? AND inicio=?",
                    (unidade["id"], mes),
                ).fetchone()
                if (desconto_atual and fonte["desconto"] is not None and
                        desconto_atual["desconto"] != fonte["desconto"]):
                    motivos.append(f"desconto cadastrado difere da planilha em {mes}")
        if mes_fatura not in MESES:
            motivos.append("competência fora de setembro/outubro")
        if motivos:
            pendencias.append({"fatura_id": documento["id"], "uc": numero,
                               "motivos": sorted(set(motivos)), "mapa_geral": geral})
            continue
        if principal["usina"] not in usinas or usinas[principal["usina"]] not in modalidades:
            pendencias.append({"fatura_id": documento["id"], "uc": numero,
                               "motivos": ["usina sem modalidade confirmada"]})
            continue
        acoes.append({"fatura_id": documento["id"], "uc": numero,
                      "competencia": mes_fatura, "unidade": unidade,
                      "usina": principal["usina"], "usina_id": usinas[principal["usina"]],
                      "fontes": por_mes, "mapa_geral": geral})
    return {"acoes": acoes, "pendencias": pendencias, "hashes": hashes,
            "usinas": usinas, "modalidades": modalidades,
            "baseline": [dict(r) for r in banco.execute(
                "SELECT id,usina_id,versao FROM faturamento_unidades WHERE usina_id IS NOT NULL ORDER BY id")],
            "documentos": len(documentos)}


def hash_plano(plano):
    return hashlib.sha256(json.dumps(plano, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def aplicar(banco, plano, arquivo_banco, hash_esperado, fontes_lidas):
    pasta_backup = arquivo_banco.parent / "backups"
    pasta_backup.mkdir(exist_ok=True)
    destino = pasta_backup / ("ecosol-antes-reconciliar-faturas-usinas-" +
                              datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ") + ".db")
    copia = sqlite3.connect(destino)
    try:
        banco.backup(copia)
    finally:
        copia.close()
    resumo = Counter()
    banco.execute("BEGIN IMMEDIATE")
    try:
        if hash_plano(planejar(banco, *fontes_lidas)) != hash_esperado:
            raise RuntimeError("Prévia mudou desde a confirmação; nenhuma alteração aplicada")
        if banco.execute("SELECT 1 FROM faturamento_vinculos_usina LIMIT 1").fetchone():
            raise RuntimeError("Histórico de usinas mudou; nenhuma alteração aplicada")
        # Antes da nova evidência de setembro, as vinculações correntes eram
        # conhecidas somente desde outubro. Isso impede retroagir transferências.
        banco.executemany(
            "INSERT INTO faturamento_vinculos_usina(unidade_id,inicio,usina_id) VALUES(?, '2026-10', ?)",
            [(x["id"], x["usina_id"]) for x in plano["baseline"]],
        )
        resumo["historicos_outubro"] = len(plano["baseline"])
        for usina_id, modalidade in plano["modalidades"].items():
            if plano["usinas"].get("ecosol") == usina_id:
                continue
            banco.execute("INSERT INTO usinas_modalidades(usina_id,inicio,modalidade) VALUES(?, '2026-09', ?)",
                          (usina_id, modalidade))
            resumo["modalidades_setembro"] += 1
        provisorios = {}
        for item in plano["acoes"]:
            unidade = item["unidade"]
            fonte = item["fontes"].get("2026-10") or item["fontes"]["2026-09"]
            if unidade is None:
                chave = (item["usina"], normalizar(fonte["nome"]))
                cliente_id = provisorios.get(chave)
                if cliente_id is None:
                    cliente_id = banco.execute(
                        "INSERT INTO faturamento_clientes_provisorios(nome,chave_origem,origem) VALUES(?,?,?)",
                        (fonte["nome"], "reconciliacao-faturas-2026:" + hashlib.sha256(
                            repr(chave).encode()).hexdigest(), "Planilha da usina setembro/outubro 2026"),
                    ).lastrowid
                    provisorios[chave] = cliente_id
                    resumo["clientes_provisorios"] += 1
                unidade_id = banco.execute("""
                    INSERT INTO faturamento_unidades
                      (uc,uc_antiga,nome,documento,email,dia_vencimento,usina_id,cliente_provisorio_id)
                    VALUES(?,?,?,?,?,?,?,?)
                """, (item["uc"], fonte["uc_antiga"] or None, fonte["nome"], "", "", 0,
                      item["usina_id"], cliente_id)).lastrowid
                banco.execute("INSERT INTO faturamento_vinculos_usina(unidade_id,inicio,usina_id) VALUES(?, '2026-10', ?)",
                              (unidade_id, item["usina_id"]))
                resumo["ucs_criadas"] += 1
            else:
                unidade_id = unidade["id"]
                if unidade["usina_id"] != item["usina_id"]:
                    banco.execute("UPDATE faturamento_unidades SET usina_id=?,versao=versao+1 WHERE id=?",
                                  (item["usina_id"], unidade_id))
                    if unidade["usina_id"] is None:
                        banco.execute("INSERT INTO faturamento_vinculos_usina(unidade_id,inicio,usina_id) VALUES(?, '2026-10', ?)",
                                      (unidade_id, item["usina_id"]))
                    else:
                        banco.execute("UPDATE faturamento_vinculos_usina SET usina_id=?,origem_usina_id=? WHERE unidade_id=? AND inicio='2026-10'",
                                      (item["usina_id"], unidade["usina_id"], unidade_id))
                    resumo["ucs_vinculadas"] += 1
            setembro = item["fontes"].get("2026-09")
            if setembro:
                banco.execute("INSERT INTO faturamento_vinculos_usina(unidade_id,inicio,usina_id) VALUES(?, '2026-09', ?)",
                              (unidade_id, item["usina_id"]))
                resumo["vinculos_setembro"] += 1
                desconto = setembro["desconto"]
                anterior = banco.execute("SELECT desconto FROM faturamento_descontos_uc WHERE unidade_id=? AND inicio='2026-09'",
                                         (unidade_id,)).fetchone()
                if desconto is not None and anterior is None:
                    banco.execute("INSERT INTO faturamento_descontos_uc(unidade_id,inicio,desconto) VALUES(?, '2026-09', ?)",
                                  (unidade_id, desconto))
                    resumo["descontos_setembro"] += 1
            outubro = item["fontes"].get("2026-10")
            if outubro and outubro["desconto"] is not None:
                anterior = banco.execute("SELECT desconto FROM faturamento_descontos_uc WHERE unidade_id=? AND inicio='2026-10'",
                                         (unidade_id,)).fetchone()
                if anterior is None:
                    banco.execute("INSERT INTO faturamento_descontos_uc(unidade_id,inicio,desconto) VALUES(?, '2026-10', ?)",
                                  (unidade_id, outubro["desconto"]))
                    resumo["descontos_outubro"] += 1
            banco.execute("INSERT INTO faturamento_historico(unidade_id,responsavel,acao,dados) VALUES(?,?,?,?)",
                          (unidade_id, "Reconciliação autorizada pelo usuário", "UC reconciliada com planilhas da usina",
                           json.dumps({"hash_previa": hash_esperado, "fatura_id": item["fatura_id"],
                                       "uc": item["uc"], "usina": item["usina"],
                                       "fontes": item["fontes"], "mapa_geral": item["mapa_geral"]},
                                      ensure_ascii=False)))
        if banco.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("Falha na integridade do banco")
        banco.commit()
    except Exception:
        banco.rollback()
        raise
    return {"backup": str(destino), "resultado": dict(resumo)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("planilhas", type=Path)
    parser.add_argument("banco", type=Path)
    parser.add_argument("--apply", metavar="HASH")
    args = parser.parse_args()
    fontes_lidas = fontes(args.planilhas)
    banco = sqlite3.connect(args.banco if args.apply else f"file:{args.banco}?mode=ro",
                            uri=not args.apply, timeout=15)
    banco.row_factory = sqlite3.Row
    banco.execute("PRAGMA foreign_keys=ON")
    try:
        plano = planejar(banco, *fontes_lidas)
        chave = hash_plano(plano)
        resultado = {"hash_previa": chave, "faturas": plano["documentos"],
                     "elegiveis": len(plano["acoes"]), "pendentes": len(plano["pendencias"]),
                     "ucs_a_criar": sum(x["unidade"] is None for x in plano["acoes"]),
                     "ucs_a_vincular": sum(x["unidade"] is not None and x["unidade"]["usina_id"] != x["usina_id"]
                                           for x in plano["acoes"]),
                     "pendencias_por_motivo": dict(Counter(m for x in plano["pendencias"] for m in x["motivos"])),
                     "pendencias": plano["pendencias"]}
        if args.apply:
            if args.apply != chave:
                raise ValueError("Hash da prévia diferente; nenhuma alteração aplicada")
            resultado["aplicacao"] = aplicar(banco, plano, args.banco, chave, fontes_lidas)
        print(json.dumps(resultado, ensure_ascii=False, indent=2))
    finally:
        banco.close()


if __name__ == "__main__":
    main()
