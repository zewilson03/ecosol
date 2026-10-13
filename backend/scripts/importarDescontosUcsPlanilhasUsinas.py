"""Importa descontos por UC das abas setembro/outubro usadas no vínculo de clientes.

Executar sem --apply para obter a prévia. Aplicar exige o hash da prévia.
As planilhas nunca são modificadas. O conflito com desconto já cadastrado é
mantido por padrão; --conflito-planilha exige decisão humana explícita.
"""

import argparse
import hashlib
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_DOWN
from pathlib import Path

from openpyxl import load_workbook

from vincularClientesPorPlanilhasUsinas import ARQUIVOS, normalizar, uc


ORIGEM_CLIENTES = "Planilhas de setembro/outubro de 2026 por usina"
ACAO_VINCULO = "UC vinculada a cliente por planilha da usina"
APROVADA = "Aprovada — aguardando emissão"


def percentual(valor):
    if valor is None or str(valor).strip() == "":
        return None
    texto = str(valor).strip()
    tem_sinal = texto.endswith("%")
    try:
        numero = Decimal(texto.rstrip("%").replace(",", "."))
    except (InvalidOperation, ValueError):
        return None
    if not numero.is_finite() or numero < 0:
        return None
    if not tem_sinal and numero <= 1:
        numero *= 100
    if numero > 100:
        return None
    # Regra validada: desprezar todas as casas decimais do percentual,
    # sem arredondar (23,73448% -> 23%).
    return format(numero.quantize(Decimal("1"), rounding=ROUND_DOWN), ".2f")


def ler_fontes(pasta):
    linhas = {}
    por_uc_mes = defaultdict(list)
    hashes = {}
    for usina, nome_arquivo in ARQUIVOS.items():
        arquivo = pasta / nome_arquivo
        if not arquivo.is_file():
            raise ValueError(f"Planilha ausente: {nome_arquivo}")
        hashes[nome_arquivo] = hashlib.sha256(arquivo.read_bytes()).hexdigest()
        livro = load_workbook(arquivo, read_only=True, data_only=True)
        try:
            for aba in livro.worksheets:
                titulo = normalizar(aba.title)
                mes = "2026-10" if re.search(r"\bout(ubro)?\b", titulo) else (
                    "2026-09" if re.search(r"\bset(embro)?\b", titulo) else None
                )
                if mes is None:
                    continue
                for numero_linha, valores in enumerate(aba.iter_rows(min_row=2, values_only=True), 2):
                    numero_uc = uc(valores[3] if len(valores) > 3 else None)
                    if not numero_uc:
                        continue
                    desconto = percentual(valores[6] if len(valores) > 6 else None)
                    chave_linha = (nome_arquivo, aba.title, numero_linha)
                    linhas[chave_linha] = (numero_uc, desconto, mes, usina)
                    por_uc_mes[(usina, numero_uc, mes)].append(desconto)
        finally:
            livro.close()
        if hashlib.sha256(arquivo.read_bytes()).hexdigest() != hashes[nome_arquivo]:
            raise ValueError(f"Planilha alterada durante a leitura: {nome_arquivo}")
    return linhas, por_uc_mes, hashes


def planejar(banco, fontes, por_uc_mes, hashes, corrigir_conflito):
    unidades = banco.execute("""
        SELECT u.id,u.uc,u.usina_id,u.cliente_provisorio_id,u.versao,s.nome AS usina
        FROM faturamento_unidades u
        JOIN usinas s ON s.id=u.usina_id
        JOIN faturamento_clientes_provisorios p ON p.id=u.cliente_provisorio_id
        WHERE p.origem=? ORDER BY u.id
    """, (ORIGEM_CLIENTES,)).fetchall()
    if len(unidades) != 576:
        raise ValueError("O conjunto de 576 UCs mudou; revise o escopo antes de importar.")
    historicos = {}
    for row in banco.execute("SELECT unidade_id,dados FROM faturamento_historico WHERE acao=? ORDER BY id DESC", (ACAO_VINCULO,)):
        historicos.setdefault(row["unidade_id"], json.loads(row["dados"]))
    acoes, pendencias = [], []
    for unidade in unidades:
        numero_uc = uc(unidade["uc"])
        vinculo = historicos.get(unidade["id"])
        if not vinculo:
            pendencias.append({"uc": numero_uc, "motivo": "Vínculo sem origem auditável"})
            continue
        arquivo, aba, linha = vinculo.get("arquivo"), vinculo.get("aba"), vinculo.get("linha")
        if ARQUIVOS.get(unidade["usina"]) != arquivo:
            pendencias.append({"uc": numero_uc, "motivo": "Usina ou arquivo alterado"})
            continue
        fonte = fontes.get((arquivo, aba, linha))
        if not fonte or (fonte[0], fonte[3]) != (numero_uc, unidade["usina"]):
            pendencias.append({"uc": numero_uc, "motivo": "UC não confere com a linha de origem"})
            continue
        _, desconto, mes, _ = fonte
        if mes != vinculo.get("competencia"):
            pendencias.append({"uc": numero_uc, "motivo": "Competência não confere com o vínculo"})
            continue
        valores = por_uc_mes[(unidade["usina"], numero_uc, mes)]
        if desconto is None or any(valor != desconto for valor in valores):
            pendencias.append({"uc": numero_uc, "motivo": "Percentual ausente, inválido ou divergente"})
            continue
        atual = banco.execute(
            "SELECT id,inicio,desconto FROM faturamento_descontos_uc WHERE unidade_id=? AND inicio<=? ORDER BY inicio DESC LIMIT 1",
            (unidade["id"], mes),
        ).fetchone()
        if atual and Decimal(atual["desconto"]) == Decimal(desconto):
            tipo = "ja_cadastrado"
        elif atual and atual["inicio"] == mes:
            tipo = "corrigir" if corrigir_conflito and numero_uc == "375491101262" else "conflito"
        else:
            tipo = "inserir"
        if tipo in {"inserir", "corrigir"} and banco.execute(
            "SELECT 1 FROM faturamento_documentos WHERE unidade_id=? AND competencia>=? AND status=? LIMIT 1",
            (unidade["id"], mes, APROVADA),
        ).fetchone():
            tipo = "cobranca_aprovada"
        if tipo in {"conflito", "cobranca_aprovada"}:
            pendencias.append({"uc": numero_uc, "motivo": tipo, "planilha": desconto,
                               "cadastrado": atual["desconto"] if atual else None})
            continue
        acoes.append({"tipo": tipo, "unidade_id": unidade["id"], "uc": numero_uc,
                      "usina_id": unidade["usina_id"], "cliente_provisorio_id": unidade["cliente_provisorio_id"],
                      "versao": unidade["versao"], "inicio": mes, "desconto": desconto,
                      "anterior_id": atual["id"] if atual else None,
                      "anterior_inicio": atual["inicio"] if atual else None,
                      "anterior_desconto": atual["desconto"] if atual else None,
                      "arquivo": arquivo, "aba": aba, "linha": linha, "hash_arquivo": hashes[arquivo]})
    plano = {"acoes": acoes, "pendencias": pendencias, "hashes": hashes,
             "corrigir_conflito": corrigir_conflito}
    hash_plano = hashlib.sha256(json.dumps(plano, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return plano, hash_plano


def aplicar(banco, plano, banco_arquivo, hash_plano):
    pasta = banco_arquivo.parent / "backups"
    pasta.mkdir(exist_ok=True)
    data = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%S-%fZ")
    destino = pasta / f"ecosol-antes-descontos-usinas-{data}.db"
    copia = sqlite3.connect(destino)
    try:
        banco.backup(copia)
    finally:
        copia.close()
    contagem = Counter()
    banco.execute("BEGIN IMMEDIATE")
    try:
        for item in plano["acoes"]:
            if item["tipo"] == "ja_cadastrado":
                contagem["ja_cadastrado"] += 1
                continue
            unidade = banco.execute(
                "SELECT usina_id,cliente_provisorio_id,versao FROM faturamento_unidades WHERE id=?",
                (item["unidade_id"],),
            ).fetchone()
            if not unidade or (unidade["usina_id"], unidade["cliente_provisorio_id"], unidade["versao"]) != (
                item["usina_id"], item["cliente_provisorio_id"], item["versao"]
            ):
                raise RuntimeError("UC alterada desde a prévia; transação cancelada")
            atual = banco.execute(
                "SELECT id,inicio,desconto FROM faturamento_descontos_uc WHERE unidade_id=? AND inicio<=? ORDER BY inicio DESC LIMIT 1",
                (item["unidade_id"], item["inicio"]),
            ).fetchone()
            estado_atual = (atual["id"], atual["inicio"], atual["desconto"]) if atual else (None, None, None)
            if estado_atual != (
                item["anterior_id"], item["anterior_inicio"], item["anterior_desconto"]
            ):
                raise RuntimeError("Desconto alterado desde a prévia; transação cancelada")
            if banco.execute(
                "SELECT 1 FROM faturamento_documentos WHERE unidade_id=? AND competencia>=? AND status=? LIMIT 1",
                (item["unidade_id"], item["inicio"], APROVADA),
            ).fetchone():
                raise RuntimeError("Cobrança aprovada desde a prévia; transação cancelada")
            if item["tipo"] == "inserir":
                banco.execute(
                    "INSERT INTO faturamento_descontos_uc(unidade_id,inicio,desconto) VALUES(?,?,?)",
                    (item["unidade_id"], item["inicio"], item["desconto"]),
                )
                acao = "Desconto da UC importado da planilha da usina"
            elif item["tipo"] == "corrigir":
                r = banco.execute(
                    "UPDATE faturamento_descontos_uc SET desconto=? WHERE id=? AND desconto=?",
                    (item["desconto"], item["anterior_id"], item["anterior_desconto"]),
                )
                if r.rowcount != 1:
                    raise RuntimeError("Desconto mudou durante a correção; transação cancelada")
                acao = "Desconto da UC corrigido pela planilha da usina"
            else:
                raise RuntimeError("Ação de desconto inesperada")
            banco.execute("UPDATE faturamento_unidades SET versao=versao+1 WHERE id=?", (item["unidade_id"],))
            banco.execute(
                "INSERT INTO faturamento_historico(unidade_id,responsavel,acao,dados) VALUES(?,?,?,?)",
                (item["unidade_id"], "Importação autorizada pelo usuário", acao,
                 json.dumps({"hash_previa": hash_plano, "arquivo": item["arquivo"],
                             "hash_arquivo": item["hash_arquivo"], "aba": item["aba"],
                             "linha": item["linha"], "competencia": item["inicio"],
                             "desconto": item["desconto"], "anterior": item["anterior_desconto"]},
                            ensure_ascii=False)),
            )
            contagem[item["tipo"]] += 1
        banco.commit()
    except Exception:
        banco.rollback()
        raise
    return {"backup": str(destino), **contagem}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--planilhas", type=Path, required=True)
    parser.add_argument("--banco", type=Path, required=True)
    parser.add_argument("--apply", metavar="HASH")
    parser.add_argument("--conflito-planilha", action="store_true")
    args = parser.parse_args()
    if not args.planilhas.is_dir() or not args.banco.is_file():
        parser.error("Pasta de planilhas ou banco ausente")
    fontes, por_uc_mes, hashes = ler_fontes(args.planilhas)
    banco = sqlite3.connect(args.banco if args.apply else f"file:{args.banco}?mode=ro",
                            uri=not args.apply, timeout=15)
    banco.row_factory = sqlite3.Row
    banco.execute("PRAGMA foreign_keys=ON")
    try:
        plano, hash_plano = planejar(banco, fontes, por_uc_mes, hashes, args.conflito_planilha)
        resumo = {"hash_previa": hash_plano, "acoes": dict(Counter(a["tipo"] for a in plano["acoes"])),
                  "pendencias": dict(Counter(p["motivo"] for p in plano["pendencias"])),
                  "competencias": dict(Counter(a["inicio"] for a in plano["acoes"])),
                  "corrigir_conflito": args.conflito_planilha}
        if args.apply:
            if args.apply != hash_plano:
                raise ValueError("A prévia mudou; nada foi gravado")
            resumo["resultado"] = aplicar(banco, plano, args.banco, hash_plano)
        print(json.dumps(resumo, ensure_ascii=False, indent=2))
    finally:
        banco.close()


if __name__ == "__main__":
    main()
