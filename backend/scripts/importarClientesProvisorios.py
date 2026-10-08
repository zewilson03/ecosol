"""Importa cadastros provisórios de teste a partir de outubro e setembro.

Uso: python importarClientesProvisorios.py --planilhas PASTA --banco ARQUIVO
     [--aplicar]

Sem --aplicar, apenas confere e resume. Não altera as planilhas.
"""

import argparse
import hashlib
import json
import re
import sqlite3
import unicodedata
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def texto_normalizado(valor):
    sem_acentos = "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", str(valor or "").lower())
        if unicodedata.category(caractere) != "Mn"
    )
    return " ".join(sem_acentos.split())


def uc_normalizada(valor):
    return re.sub(r"\D", "", str(valor or "")).lstrip("0")


def mes_da_aba(nome):
    nome = texto_normalizado(nome)
    if "out" in nome:
        return "2026-10"
    if "set" in nome:
        return "2026-09"
    return None


def desconto_percentual(valor):
    try:
        numero = Decimal(str(valor).replace(",", "."))
    except (InvalidOperation, ValueError):
        return None
    if 0 <= numero <= 1:
        numero *= 100
    if numero < 0 or numero > 100 or numero.as_tuple().exponent < -5:
        return None
    arredondado = numero.quantize(Decimal("0.01"))
    return format(arredondado, "f") if arredondado == numero else None


def dia_preferido(observacoes):
    dias = {
        int(match)
        for match in re.findall(
            r"(?:vencimento|pagamento)\s+(?:sempre\s+|para\s+o\s+|todo\s+)?dia\s+(\d{1,2})\b",
            texto_normalizado(observacoes),
        )
        if 1 <= int(match) <= 31
    }
    return next(iter(dias)) if len(dias) == 1 else 0


def faturas_ativas(banco):
    ucs = defaultdict(list)
    for item in banco.execute(
        "SELECT id,nome,dados,extracao FROM faturamento_documentos "
        "WHERE classificacao='fatura' AND excluido_em IS NULL"
    ):
        dados = json.loads(item["dados"])
        extracao = json.loads(item["extracao"])
        uc_salva = uc_normalizada(dados.get("uc"))
        uc_extraida = uc_normalizada(extracao.get("dados", {}).get("uc"))
        nome_pdf = str(item["nome"])
        numero_arquivo = uc_normalizada(
            nome_pdf[:-4] if nome_pdf.lower().endswith(".pdf") else nome_pdf
        )
        # Erro conhecido de preenchimento: o número do arquivo foi gravado como UC.
        # Usar a UC extraída apenas para localizar o cadastro; não editar a fatura.
        uc = (
            uc_extraida
            if uc_salva == numero_arquivo and uc_extraida and uc_extraida != uc_salva
            else uc_salva or uc_extraida
        )
        if uc:
            ucs[uc].append(item["id"])
    return ucs


def procurar_linhas(pasta, ucs):
    resultados = defaultdict(lambda: {"2026-10": [], "2026-09": []})
    for arquivo in sorted(pasta.rglob("*.xlsx")):
        planilha = load_workbook(arquivo, read_only=True, data_only=True)
        try:
            for aba in planilha.worksheets:
                mes = mes_da_aba(aba.title)
                if mes is None:
                    continue
                cabecalho = [
                    texto_normalizado(valor)
                    for valor in next(aba.iter_rows(max_row=1, values_only=True))
                ]
                coluna_uc = next(
                    (
                        indice
                        for indice, valor in enumerate(cabecalho)
                        if "nova uc" in valor or valor == "uc2"
                    ),
                    3,
                )
                for linha, valores in enumerate(
                    aba.iter_rows(min_row=2, values_only=True), 2
                ):
                    uc = uc_normalizada(
                        valores[coluna_uc] if len(valores) > coluna_uc else ""
                    )
                    if uc not in ucs:
                        continue
                    resultados[uc][mes].append(
                        {
                            "arquivo": str(arquivo),
                            "aba": aba.title,
                            "linha": linha,
                            "nome": str(valores[1] or "").strip(),
                            "desconto": desconto_percentual(valores[6]),
                            "dia": dia_preferido(
                                " ".join(str(valor or "") for valor in valores[18:])
                            ),
                        }
                    )
        finally:
            planilha.close()
    return resultados


def preparar_candidatos(ucs, encontrados):
    candidatos = []
    pendencias = []
    for uc in sorted(ucs):
        meses = encontrados.get(uc)
        if not meses:
            pendencias.append({"uc": uc, "motivo": "UC não localizada"})
            continue
        mes = "2026-10" if meses["2026-10"] else "2026-09"
        linhas = meses[mes]
        nomes = {texto_normalizado(linha["nome"]) for linha in linhas}
        descontos = {linha["desconto"] for linha in linhas}
        if len(nomes) != 1 or not nomes.pop() or len(descontos) != 1:
            pendencias.append({"uc": uc, "motivo": "Linhas divergentes"})
            continue
        if None in descontos or len(linhas[0]["nome"]) > 120:
            pendencias.append({"uc": uc, "motivo": "Nome ou desconto inválido"})
            continue
        dias = {linha["dia"] for linha in linhas if linha["dia"]}
        if len(dias) > 1:
            pendencias.append({"uc": uc, "motivo": "Dia preferido divergente"})
            continue
        candidato = {
            **linhas[0],
            "uc": uc,
            "mes": mes,
            "dia": next(iter(dias)) if dias else 0,
            "faturas": ucs[uc],
        }
        candidatos.append(candidato)
    return candidatos, pendencias


def aplicar(banco, candidatos):
    for tabela, coluna in (
        ("faturamento_unidades", "cliente_provisorio_id"),
        ("faturamento_unidades", "cliente_id"),
    ):
        if coluna not in {
            item["name"] for item in banco.execute(f"PRAGMA table_info({tabela})")
        }:
            raise RuntimeError("Execute a migração do Ecosol antes da importação.")
    if banco.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' "
        "AND name='faturamento_clientes_provisorios'"
    ).fetchone() is None:
        raise RuntimeError("Tabela de clientes provisórios ainda não criada.")
    criados = 0
    vinculados = 0
    ignorados = []
    banco.execute("BEGIN IMMEDIATE")
    try:
        for item in candidatos:
            existente = banco.execute(
                "SELECT id FROM faturamento_unidades WHERE uc=?", (item["uc"],)
            ).fetchone()
            if existente:
                ignorados.append({"uc": item["uc"], "motivo": "UC já cadastrada"})
                continue
            chave = hashlib.sha256(
                ("teste-2026:" + texto_normalizado(item["nome"])).encode("utf-8")
            ).hexdigest()
            cliente = banco.execute(
                "SELECT id FROM faturamento_clientes_provisorios WHERE chave_origem=?",
                (chave,),
            ).fetchone()
            if cliente is None:
                cliente_id = banco.execute(
                    "INSERT INTO faturamento_clientes_provisorios"
                    "(nome,chave_origem,origem) VALUES (?,?,?)",
                    (item["nome"], chave, "Planilhas de outubro/setembro de 2026 — teste"),
                ).lastrowid
                criados += 1
            else:
                cliente_id = cliente["id"]
            unidade_id = banco.execute(
                "INSERT INTO faturamento_unidades"
                "(uc,nome,documento,email,dia_vencimento,cliente_provisorio_id) "
                "VALUES (?,?,?,?,?,?)",
                (item["uc"], item["nome"], "", "", item["dia"], cliente_id),
            ).lastrowid
            banco.execute(
                "INSERT INTO faturamento_descontos_uc(unidade_id,inicio,desconto) "
                "VALUES (?,?,?)",
                (unidade_id, item["mes"], item["desconto"]),
            )
            banco.execute(
                "INSERT INTO faturamento_historico"
                "(unidade_id,responsavel,acao,dados) VALUES (?,?,?,?)",
                (
                    unidade_id,
                    "Importação de teste",
                    "Cadastro provisório de cliente e UC",
                    json.dumps(
                        {
                            "arquivo": Path(item["arquivo"]).name,
                            "aba": item["aba"],
                            "linha": item["linha"],
                            "competencia_desconto": item["mes"],
                            "desconto": item["desconto"],
                            "dia_preferido_confirmado": item["dia"] > 0,
                        },
                        ensure_ascii=False,
                    ),
                ),
            )
            vinculados += 1
        banco.commit()
    except Exception:
        banco.rollback()
        raise
    return criados, vinculados, ignorados


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--planilhas", type=Path, required=True)
    parser.add_argument("--banco", type=Path, required=True)
    parser.add_argument("--aplicar", action="store_true")
    args = parser.parse_args()
    if not args.planilhas.is_dir() or not args.banco.is_file():
        parser.error("Pasta de planilhas ou banco não encontrado.")
    banco = sqlite3.connect(
        str(args.banco) if args.aplicar else f"file:{args.banco}?mode=ro",
        uri=not args.aplicar,
        timeout=10,
    )
    banco.row_factory = sqlite3.Row
    try:
        ucs = faturas_ativas(banco)
        candidatos, pendencias = preparar_candidatos(
            ucs, procurar_linhas(args.planilhas, ucs)
        )
        if args.aplicar:
            criados, vinculados, ignorados = aplicar(banco, candidatos)
        else:
            criados, vinculados, ignorados = 0, 0, []
        print(
            json.dumps(
                {
                    "modo": "aplicado" if args.aplicar else "conferência",
                    "ucs_faturas": len(ucs),
                    "ucs_encontradas": len(candidatos),
                    "em_outubro": sum(i["mes"] == "2026-10" for i in candidatos),
                    "em_setembro": sum(i["mes"] == "2026-09" for i in candidatos),
                    "sem_dia_preferido": sum(i["dia"] == 0 for i in candidatos),
                    "clientes_criados": criados,
                    "ucs_vinculadas": vinculados,
                    "pendencias": pendencias + ignorados,
                },
                ensure_ascii=False,
            )
        )
    finally:
        banco.close()


if __name__ == "__main__":
    main()
