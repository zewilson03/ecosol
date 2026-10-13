"""Consolida clientes provisórios por nome, UC e usina em setembro/outubro.

Leitura das planilhas é estritamente read-only. Sem --apply HASH, apenas gera
uma prévia. A aplicação exige o hash da prévia e cria backup do banco.
"""

import argparse
import hashlib
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ARQUIVOS = {
    "2000 Solar": "24 - Venda de Energia - Usina 2000 Solar.xlsx",
    "Abadiânia": "01 - Venda de Energia - Usina Abadiania 2026.xlsx",
    "CLAUDIO": "13 - Venda de Energia - Usina Claudio 2026.xlsx",
    "MÁRCIO": "06 - Venda de Energia - Usina Marcio 2026.xlsx",
    "NEOVOLTS": "18- Venda de Energia - Usina Neovolts 2026.xlsx",
    "OURO VERDE": "02 - Venda de Energia - Usina Ouro Verde 2026.xlsx",
    "ROYER": "20 - Venda de Energia - Usina Royer GO 2026.xlsx",
    "SAN GERMAN": "04 - Venda de Energia - Usina San German  Participaçoes Societarias Ltda 2026.xlsx.xlsx",
    "Solar Green": "10 - Venda de Energia - Usina Solar Green 2026.xlsx",
    "São Francisco": "03 - Venda de Energia - Usina São Francisco 2026.xlsx.xlsx",
}


def normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.casefold().split())


def uc(valor):
    return re.sub(r"\D", "", str(valor or "")).lstrip("0")


DOCUMENTO = re.compile(
    r"(?<!\d)(?:\d{3}\.?\d{3}\.?\d{3}[-.]?\d{2}|"
    r"\d{2}\.?\d{3}\.?\d{3}[/.-]?\d{4}[-.]?\d{2})(?!\d)"
)


def nome_valido(valor):
    nome = " ".join(str(valor or "").split())
    nome = DOCUMENTO.sub("", nome)
    # Algumas células trazem CNPJ digitado com pontuação incorreta. Ainda
    # assim é uma sequência documental longa ao fim, não parte do nome.
    nome = re.sub(r"\s*[-–—]?\s*\d[\d./-]{10,28}$", "", nome)
    nome = re.sub(r"\b(?:CPF|CNPJ)\b\s*[:\-]?", "", nome, flags=re.IGNORECASE)
    nome = nome.strip(" .-/–—")
    nome = " ".join(nome.split())
    if not nome or len(nome) > 120 or normalizar(nome) in {"nan", "-", "0"}:
        return None
    return nome


def nome_comparavel(valor):
    # Os cadastros antigos frequentemente embutem CPF/CNPJ ao final do nome.
    # A informação documental é ignorada na comparação e nunca importada.
    nome = normalizar(nome_valido(valor))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", nome).split())


def dados_planilhas(pasta):
    linhas = defaultdict(lambda: defaultdict(list))
    fontes = {}
    for usina, nome_arquivo in ARQUIVOS.items():
        arquivo = pasta / nome_arquivo
        if not arquivo.is_file():
            raise ValueError(f"Planilha obrigatória ausente: {nome_arquivo}")
        fontes[usina] = hashlib.sha256(arquivo.read_bytes()).hexdigest()
        planilha = load_workbook(arquivo, read_only=True, data_only=True)
        try:
            meses = set()
            for aba in planilha.worksheets:
                titulo = normalizar(aba.title)
                mes = "2026-10" if re.search(r"\bout(ubro)?\b", titulo) else (
                    "2026-09" if re.search(r"\bset(embro)?\b", titulo) else None
                )
                if mes is None:
                    continue
                if mes in meses:
                    raise ValueError(f"Abas duplicadas para {mes} em {nome_arquivo}")
                meses.add(mes)
                cabecalho = next(aba.iter_rows(max_row=1, values_only=True))
                if normalizar(cabecalho[1]) not in {"clientes", "cliente"}:
                    raise ValueError(f"Coluna de cliente inesperada em {nome_arquivo}/{aba.title}")
                if not any(p in normalizar(cabecalho[3]) for p in ("uc2", "nova uc", "uc nova")):
                    raise ValueError(f"Coluna de UC nova inesperada em {nome_arquivo}/{aba.title}")
                for numero, valores in enumerate(aba.iter_rows(min_row=2, values_only=True), 2):
                    chave_uc = uc(valores[3] if len(valores) > 3 else None)
                    if not chave_uc:
                        continue
                    nome = nome_valido(valores[1] if len(valores) > 1 else None)
                    linhas[(usina, chave_uc)][mes].append({
                        "nome": nome, "arquivo": nome_arquivo, "aba": aba.title,
                        "linha": numero,
                    })
            if meses != {"2026-09", "2026-10"}:
                raise ValueError(f"Setembro/outubro não encontrados em {nome_arquivo}")
        finally:
            planilha.close()
        if hashlib.sha256(arquivo.read_bytes()).hexdigest() != fontes[usina]:
            raise ValueError(f"Planilha alterada durante a leitura: {nome_arquivo}")
    return linhas, fontes


def planejar(banco, planilhas, fontes):
    usinas = {r["nome"]: r["id"] for r in banco.execute("SELECT id,nome FROM usinas")}
    if set(ARQUIVOS) - set(usinas):
        raise ValueError("O catálogo de usinas mudou; atualize a prévia.")
    unidades = banco.execute("""
        SELECT u.id,u.uc,u.nome,u.cliente_id,u.cliente_provisorio_id,u.usina_id,
               u.versao,p.nome AS nome_provisorio,p.origem AS origem_provisorio
        FROM faturamento_unidades u
        LEFT JOIN faturamento_clientes_provisorios p ON p.id=u.cliente_provisorio_id
    """).fetchall()
    por_uc = defaultdict(list)
    for (usina, numero), meses in planilhas.items():
        por_uc[numero].append((usina, set(meses)))
    acoes, pendencias = [], []
    resumo = Counter()
    for unidade in unidades:
        usina = next((nome for nome, ident in usinas.items() if ident == unidade["usina_id"]), None)
        if usina not in ARQUIVOS:
            continue
        numero = uc(unidade["uc"])
        meses = planilhas.get((usina, numero))
        if not meses:
            pendencias.append({"uc": numero, "usina": usina, "motivo": "UC não localizada na própria usina"})
            continue
        mes = "2026-10" if meses.get("2026-10") else "2026-09"
        candidatas = meses[mes]
        nomes = {normalizar(item["nome"]) for item in candidatas}
        if mes == "2026-10" and nomes == {""} and meses.get("2026-09"):
            mes = "2026-09"
            candidatas = meses[mes]
            nomes = {normalizar(item["nome"]) for item in candidatas}
        if any(outra != usina and (mes in competencias or
               (mes == "2026-09" and "2026-10" in competencias))
               for outra, competencias in por_uc[numero]):
            pendencias.append({"uc": numero, "usina": usina, "motivo": "UC presente em mais de uma usina na competência relevante"})
            continue
        if len(nomes) != 1 or not next(iter(nomes)):
            pendencias.append({"uc": numero, "usina": usina, "motivo": "Nome ausente ou divergente na competência"})
            continue
        nome = candidatas[0]["nome"]
        if unidade["cliente_id"] is not None:
            pendencias.append({"uc": numero, "usina": usina, "motivo": "Cliente com acesso já associado"})
            continue
        if not unidade["cliente_provisorio_id"]:
            pendencias.append({"uc": numero, "usina": usina, "motivo": "UC sem cadastro provisório"})
            continue
        acoes.append({
            "unidade_id": unidade["id"], "uc": numero, "usina": usina,
            "nome": nome, "chave_nome": normalizar(nome), "mes": mes,
            "origem": candidatas[0], "cliente_anterior": unidade["cliente_provisorio_id"],
            "nome_anterior": unidade["nome"],
            "nome_divergente": nome_comparavel(unidade["nome_provisorio"] or unidade["nome"]) != nome_comparavel(nome),
            "usina_id": unidade["usina_id"],
            "versao": unidade["versao"],
        })
        resumo[mes] += 1
    # Só consolidar identidades homônimas dentro da mesma usina. Uma pessoa
    # homônima em outra usina não é confirmada sem CPF/documento.
    grupos = defaultdict(list)
    for acao in acoes:
        grupos[(acao["usina"], acao["chave_nome"])].append(acao)
    plano = {
        "fontes": fontes,
        "acoes": sorted(acoes, key=lambda x: x["unidade_id"]),
        "pendencias": sorted(pendencias, key=lambda x: (x["usina"], x["uc"])),
        "grupos": len(grupos),
        "resumo": dict(resumo),
    }
    hash_plano = hashlib.sha256(json.dumps(plano, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return plano, grupos, hash_plano


def aplicar(banco, plano, grupos, caminho_banco, hash_plano):
    pasta_backup = caminho_banco.parent / "backups"
    pasta_backup.mkdir(exist_ok=True)
    data = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%S-%fZ")
    destino = pasta_backup / f"ecosol-antes-vincular-clientes-usinas-{data}.db"
    copia = sqlite3.connect(destino)
    try:
        banco.backup(copia)
    finally:
        copia.close()
    alteradas = 0
    criados = 0
    banco.execute("BEGIN IMMEDIATE")
    try:
        for (usina, chave_nome), acoes in sorted(grupos.items()):
            nome = next((a["nome"] for a in acoes if a["mes"] == "2026-10"), acoes[0]["nome"])
            chave = hashlib.sha256(f"usina-planilha-2026:{usina}:{chave_nome}".encode()).hexdigest()
            cliente = banco.execute(
                "SELECT id,nome FROM faturamento_clientes_provisorios WHERE chave_origem=?", (chave,)
            ).fetchone()
            if cliente is None:
                cliente_id = banco.execute(
                    "INSERT INTO faturamento_clientes_provisorios(nome,chave_origem,origem) VALUES(?,?,?)",
                    (nome, chave, "Planilhas de setembro/outubro de 2026 por usina"),
                ).lastrowid
                criados += 1
            else:
                cliente_id = cliente["id"]
                if normalizar(cliente["nome"]) != chave_nome:
                    raise RuntimeError("Cadastro consolidado existente mudou de nome")
            for acao in acoes:
                unidade = banco.execute(
                    "SELECT nome,cliente_id,cliente_provisorio_id,usina_id,versao FROM faturamento_unidades WHERE id=?",
                    (acao["unidade_id"],),
                ).fetchone()
                if (unidade["cliente_id"] is not None or unidade["cliente_provisorio_id"] != acao["cliente_anterior"]
                        or unidade["versao"] != acao["versao"] or unidade["nome"] != acao["nome_anterior"]
                        or unidade["usina_id"] != acao["usina_id"]):
                    raise RuntimeError("UC alterada desde a prévia")
                if unidade["cliente_provisorio_id"] == cliente_id and unidade["nome"] == acao["nome"]:
                    continue
                banco.execute(
                    "UPDATE faturamento_unidades SET nome=?,cliente_provisorio_id=?,versao=versao+1 WHERE id=?",
                    (acao["nome"], cliente_id, acao["unidade_id"]),
                )
                banco.execute(
                    "INSERT INTO faturamento_historico(unidade_id,responsavel,acao,dados) VALUES(?,?,?,?)",
                    (acao["unidade_id"], "Vínculo autorizado pelo usuário", "UC vinculada a cliente por planilha da usina",
                     json.dumps({"hash_previa": hash_plano, "arquivo": acao["origem"]["arquivo"],
                                 "aba": acao["origem"]["aba"], "linha": acao["origem"]["linha"],
                                 "competencia": acao["mes"], "nome_divergente": acao["nome_divergente"],
                                 "cliente_anterior": acao["cliente_anterior"], "cliente_novo": cliente_id}, ensure_ascii=False)),
                )
                alteradas += 1
        banco.commit()
    except Exception:
        banco.rollback()
        raise
    return {"backup": str(destino), "clientes_criados": criados, "ucs_vinculadas": alteradas}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--planilhas", required=True, type=Path)
    parser.add_argument("--banco", required=True, type=Path)
    parser.add_argument("--apply", metavar="HASH")
    args = parser.parse_args()
    if not args.planilhas.is_dir() or not args.banco.is_file():
        parser.error("Pasta de planilhas ou banco não encontrado")
    linhas, fontes = dados_planilhas(args.planilhas)
    banco = sqlite3.connect(args.banco if args.apply else f"file:{args.banco}?mode=ro", uri=not args.apply, timeout=15)
    banco.row_factory = sqlite3.Row
    banco.execute("PRAGMA foreign_keys=ON")
    try:
        plano, grupos, hash_plano = planejar(banco, linhas, fontes)
        resultado = {
            "hash_previa": hash_plano, "ucs_identificadas": len(plano["acoes"]),
            "grupos_de_clientes": plano["grupos"],
            "grupos_com_multiplas_ucs": sum(len(v) > 1 for v in grupos.values()),
            "nomes_diferentes_atualizados_pela_planilha": sum(a["nome_divergente"] for a in plano["acoes"]),
            "pendencias_por_motivo": dict(Counter(p["motivo"] for p in plano["pendencias"])),
            "pendencias_por_usina": dict(Counter(p["usina"] for p in plano["pendencias"])),
            "competencias": plano["resumo"],
        }
        if args.apply:
            if args.apply != hash_plano:
                raise ValueError("A prévia mudou; nenhuma alteração aplicada")
            resultado["aplicacao"] = aplicar(banco, plano, grupos, args.banco, hash_plano)
        print(json.dumps(resultado, ensure_ascii=False, indent=2))
    finally:
        banco.close()


if __name__ == "__main__":
    main()
